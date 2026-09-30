from openpyxl import load_workbook
from pathlib import Path
import csv, json, uuid

SRC=Path('Phase4_validation_results.xlsx')
OUT=Path('data'); OUT.mkdir(exist_ok=True)
NS=uuid.UUID('7c302f2a-3c0d-4f7c-a550-32c436f362fe')
wb=load_workbook(SRC,read_only=True,data_only=True)

def read(sheet):
    ws=wb[sheet]; it=ws.iter_rows(values_only=True); h=list(next(it)); return [dict(zip(h,r)) for r in it if any(x is not None for x in r)]
tees=read('Valid_Tees'); holes=read('Valid_Holes')
def uid(kind,key): return str(uuid.uuid5(NS,kind+'|'+key))
def key(r): return f"{r['course_id']}|{r['tee_name']}|{r['tee_gender']}"
# valid workbook already represents conservative gate
courses={}
for r in tees:
    cid=str(r['course_id']).strip()
    courses.setdefault(cid, {'id':uid('course',cid),'source_course_id':cid,'name':r['course_name'],'golf_union':r['union'],'source_name':'GolfRSA / SA_GolfRSA_complete','verification_status':'validated_v1'})
tee_rows=[]
for r in tees:
    k=key(r); cid=str(r['course_id']).strip()
    tee_rows.append({'id':uid('tee',k),'course_id':courses[cid]['id'],'source_course_id':cid,'tee_name':r['tee_name'],'gender_rating':r['tee_gender'],'course_rating':float(r['course_rating']),'slope_rating':int(float(r['slope_rating'])),'par':int(float(r['par'])),'source_url':r['source_file'] or ''})
teemap={key(r):uid('tee',key(r)) for r in tees}
hole_rows=[]
for r in holes:
    k=key(r)
    if k not in teemap: raise SystemExit(f'validated hole missing tee: {k}')
    hole_rows.append({'id':uid('hole',k+'|'+str(r['hole_number'])),'course_tee_id':teemap[k],'source_course_id':r['course_id'],'tee_name':r['tee_name'],'tee_gender':r['tee_gender'],'hole_number':int(r['hole_number']),'par':int(r['par']),'stroke_index':int(r['stroke_index']),'metres':int(r['distance_m']) if r['distance_m'] not in (None,'') else '','source_url':r['source_url'] or '','source_type':r['source_type'] or '','confidence':r['confidence'] or ''})
for name,rows in [('courses',courses.values()),('course_tees',tee_rows),('course_holes',hole_rows)]:
    rows=list(rows); p=OUT/f'{name}.csv';
    with p.open('w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
print(json.dumps({'courses':len(courses),'course_tees':len(tee_rows),'course_holes':len(hole_rows)},indent=2))
